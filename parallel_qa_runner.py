"""Run QA account pulls in bounded parallel lanes with a validation barrier.

The parent process is the only writer of pipeline_status.json. This avoids the
status-file and Git races that would occur if parallel child processes called
log_step.py independently.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl


BASE = Path(__file__).resolve().parent
STATUS_FILE = BASE / "pipeline_status.json"
BASELINE_FILE = BASE / "pipeline_rowcount_baseline.json"
LOG_ROOT = BASE / "pipeline_parallel_logs"
SELF_HEAL = BASE / "self_heal.py"
REQUIRED_HEADERS = {
    "QA_ID",
    "EMPLOYEE_ID",
    "QA_COACH_ID",
    "SUPERVISOR_ID",
    "Emp Name",
    "QA",
    "Immediate Supervisor",
    "LOB / Account",
}
DROP_TOLERANCE = 0.05
PUBLISH_INTERVAL_SECONDS = 60
LANE_START_STAGGER_SECONDS = 15


@dataclass(frozen=True)
class Account:
    name: str
    script: str
    directory: Path
    output: Path
    timeout_seconds: int = 1800


QUALITY = Path(r"C:\Users\Mike Woo Cerna\Documents\PB\Quality")


def account(name, folder, script, output, timeout=1800):
    directory = QUALITY / folder
    return Account(name, script, directory, directory / output, timeout)


# Static lanes are intentionally balanced using observed production runtimes.
# Skyline receives a dedicated lane; short Google-fed jobs share lane 4.
LANES = {
    "Lane 1 - Skyline": [
        account("Skyline", "Skyline", "Skyline_pull.py", "SKYLINE_RAW.xlsx", 2400),
    ],
    "Lane 2 - Hamilton": [
        account("Hamilton", "Hamilton", "Hamilton_pull.py", "HAMILTON_RAW.xlsx"),
        account("VIP", "VIP", "vip_pull.py", "VIP_RAW.xlsx"),
        account("Vermont", "Vermont", "vt_pull.py", "VT_RAW.xlsx"),
        account("Trans Iowa", "Trans Iowa", "ti_pull.py", "TI_RAW.xlsx"),
    ],
    "Lane 3 - Direct QA": [
        account("Data Carz", "Data Carz", "dc_pull.py", "DC_RAW.xlsx"),
        account("Associated Cab", "Associated Cab", "ac_pull.py", "AC_RAW.xlsx"),
        account("Ollies", "Ollies", "ol_pull.py", "OL_RAW.xlsx"),
        account("Circle Taxi", "Circle Taxi", "ct_pull.py", "CT_RAW.xlsx"),
        account("YCOV", "YCOV", "ycov_pull.py", "YCOV_RAW.xlsx"),
    ],
    "Lane 4 - Standard and Direct": [
        account("M7", "M7", "m7_pull.py", "M7_RAW.xlsx"),
        account("DMG", "DMG", "dmg_pull.py", "DMG_RAW.xlsx"),
        account("R4H", "R4H", "r4h_pull.py", "R4H_RAW.xlsx"),
        account("Parentis Health", "Parentis Health", "parentis_pull.py", "PARENTIS_RAW.xlsx"),
        account("Britelift", "Britelift", "britelift_pull.py", "BRITELIFT_RAW.xlsx"),
        account("Britelift Chat", "Britelift Chat", "britelift_pull.py", "BLC_RAW.xlsx"),
        account("RideX", "RideX", "Ridex_pull.py", "RIDEX_RAW.xlsx"),
        account("C&H", "C&H", "ch_pull.py", "CH_RAW.xlsx"),
        account("Blueline", "Blueline", "bl_pull.py", "BL_RAW.xlsx"),
        account("Reno Cab", "Reno Cab", "rc_pull.py", "RC_RAW.xlsx"),
        account("Kelowna", "Kelowna", "kel_pull.py", "KEL_RAW.xlsx"),
        account("YCDC", "YCDC", "ycdc_pull.py", "YCDC_RAW.xlsx"),
    ],
}


class StatusStore:
    def __init__(self):
        self.lock = threading.Lock()

    def _load(self):
        if STATUS_FILE.exists():
            try:
                return json.loads(STATUS_FILE.read_text(encoding="utf-8"))
            except Exception:
                pass
        now = datetime.now()
        return {
            "run_id": now.strftime("%Y-%m-%dT%H:%M:%S"),
            "started_at": now.isoformat(),
            "finished_at": None,
            "status": "running",
            "failed_at": None,
            "steps": [],
        }

    def _save(self, data):
        temp = STATUS_FILE.with_suffix(".json.tmp")
        temp.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
        os.replace(temp, STATUS_FILE)

    def set_step(self, item, status, lane, *, exit_code=None, error=None, rows=None, started_at=None):
        now = datetime.now().isoformat()
        with self.lock:
            data = self._load()
            steps = data.setdefault("steps", [])
            step = next((entry for entry in steps if entry.get("account") == item.name), None)
            if step is None:
                step = {"account": item.name, "script": item.script}
                steps.append(step)
            step.update({
                "script": item.script,
                "status": status,
                "timestamp": now,
                "lane": lane,
                "exit_code": exit_code,
                "error": error,
            })
            if started_at:
                step["started_at"] = started_at
            if rows is not None:
                step["rows"] = rows
            if status == "fail":
                data["status"] = "failed"
                data["failed_at"] = item.name
            elif data.get("status") not in {"failed", "partial"}:
                data["status"] = "running"
            self._save(data)

    def queue_all(self):
        for lane, items in LANES.items():
            for item in items:
                self.set_step(item, "queued", lane)


class LivePublisher:
    def __init__(self, enabled=True):
        self.enabled = enabled
        self.dirty = threading.Event()
        self.stop_event = threading.Event()
        self.thread = None
        self.last_publish = 0.0
        self.publish_lock = threading.Lock()

    def start(self):
        if not self.enabled:
            return
        self.thread = threading.Thread(target=self._loop, name="monitor-publisher", daemon=True)
        self.thread.start()

    def signal(self):
        if self.enabled:
            self.dirty.set()

    def _loop(self):
        while not self.stop_event.wait(3):
            if not self.dirty.is_set():
                continue
            if time.monotonic() - self.last_publish < PUBLISH_INTERVAL_SECONDS:
                continue
            self.publish("Parallel QA progress")

    def publish(self, label):
        if not self.enabled:
            return
        with self.publish_lock:
            self.dirty.clear()
            try:
                generated = subprocess.run(
                    [sys.executable, str(BASE / "generate_monitor.py")],
                    cwd=BASE,
                    capture_output=True,
                    text=True,
                    timeout=90,
                    check=False,
                )
                if generated.returncode != 0:
                    detail = (generated.stderr or generated.stdout or "unknown error").strip()
                    raise RuntimeError(f"monitor generation failed: {detail}")

                staged = subprocess.run(
                    ["git", "add", "pipeline_status.json", "pipeline_monitor.html"],
                    cwd=BASE,
                    capture_output=True,
                    timeout=30,
                    check=False,
                )
                if staged.returncode != 0:
                    raise RuntimeError("git add failed while publishing monitor progress")
                changed = subprocess.run(
                    ["git", "diff", "--cached", "--quiet"],
                    cwd=BASE,
                    capture_output=True,
                    timeout=30,
                    check=False,
                ).returncode != 0
                if changed:
                    commit = subprocess.run(
                        ["git", "commit", "-m", f"[live] {label}"],
                        cwd=BASE,
                        capture_output=True,
                        text=True,
                        timeout=45,
                        check=False,
                    )
                    if commit.returncode != 0:
                        detail = (commit.stderr or commit.stdout or "unknown error").strip()
                        raise RuntimeError(f"git commit failed: {detail}")
                    pushed = subprocess.run(
                        ["git", "push"],
                        cwd=BASE,
                        capture_output=True,
                        text=True,
                        timeout=90,
                        check=False,
                    )
                    if pushed.returncode != 0:
                        detail = (pushed.stderr or pushed.stdout or "unknown error").strip()
                        raise RuntimeError(f"git push failed: {detail}")
                self.last_publish = time.monotonic()
                print(f"[parallel] Live monitor published: {label}", flush=True)
            except Exception as exc:
                print(f"[parallel] Live monitor publish skipped: {exc}", flush=True)

    def close(self, final_label):
        if not self.enabled:
            return
        self.stop_event.set()
        if self.thread:
            self.thread.join(timeout=10)
        self.publish(final_label)


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def tail_error(path, max_lines=8):
    try:
        lines = [line.strip() for line in path.read_text(encoding="utf-8", errors="replace").splitlines() if line.strip()]
        return "\n".join(lines[-max_lines:]) or None
    except Exception:
        return None


def terminate_process_tree(process):
    if process.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(
            ["taskkill", "/PID", str(process.pid), "/T", "/F"],
            capture_output=True,
            check=False,
        )
    else:
        process.kill()


def run_account(item, lane, log_dir, store, publisher):
    started_at = datetime.now().isoformat()
    store.set_step(item, "running", lane, started_at=started_at)
    publisher.signal()
    print(f"[{lane}] RUNNING {item.name} ({item.script})", flush=True)

    log_file = log_dir / f"{slug(item.name)}.log"
    with log_file.open("w", encoding="utf-8", errors="replace") as stream:
        process = subprocess.Popen(
            [sys.executable, str(SELF_HEAL), "run-step", item.script],
            cwd=item.directory,
            stdout=stream,
            stderr=subprocess.STDOUT,
            creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0,
        )
        try:
            exit_code = process.wait(timeout=item.timeout_seconds)
        except subprocess.TimeoutExpired:
            terminate_process_tree(process)
            stream.write(f"\nTimed out after {item.timeout_seconds} seconds.\n")
            exit_code = 124

    error = tail_error(log_file) if exit_code else None
    if exit_code:
        store.set_step(item, "fail", lane, exit_code=exit_code, error=error, started_at=started_at)
        print(f"[{lane}] FAILED {item.name}; see {log_file}", flush=True)
    else:
        store.set_step(item, "checking", lane, exit_code=0, started_at=started_at)
        print(f"[{lane}] CHECKING {item.name}", flush=True)
    publisher.signal()
    return exit_code


def run_lane(lane, items, delay, log_dir, store, publisher, results):
    if delay:
        time.sleep(delay)
    for item in items:
        try:
            results[item.name] = run_account(item, lane, log_dir, store, publisher)
        except Exception as exc:
            results[item.name] = 1
            store.set_step(item, "fail", lane, exit_code=1, error=str(exc))
            publisher.signal()
            print(f"[{lane}] FAILED {item.name}: {exc}", flush=True)


def workbook_metadata(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = {cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1)) if cell.value is not None}
    rows = max(0, (ws.max_row or 1) - 1)
    wb.close()
    return rows, headers


def load_baseline():
    try:
        return json.loads(BASELINE_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def validate_outputs(results, store, publisher):
    baseline = load_baseline()
    failures = []
    for lane, items in LANES.items():
        for item in items:
            error = None
            rows = None
            if results.get(item.name, 1) != 0:
                error = "Pull process failed; validation was not attempted."
            elif not item.output.exists():
                error = f"Expected output does not exist: {item.output}"
            else:
                try:
                    rows, headers = workbook_metadata(item.output)
                    missing = sorted(REQUIRED_HEADERS - headers)
                    if rows <= 0:
                        error = "Workbook contains no data rows."
                    elif missing:
                        error = "Missing required columns: " + ", ".join(missing)
                    else:
                        prior = baseline.get(item.name)
                        if prior and rows < int(prior * (1 - DROP_TOLERANCE)):
                            error = (
                                f"Validation count drop: {prior:,} -> {rows:,} "
                                f"({(prior - rows) / prior * 100:.1f}%)."
                            )
                except Exception as exc:
                    error = f"Workbook validation failed: {exc}"

            if error:
                failures.append((item.name, error))
                store.set_step(item, "fail", lane, exit_code=1, error=error, rows=rows)
                print(f"[validation] FAIL {item.name}: {error}", flush=True)
            else:
                store.set_step(item, "pass", lane, exit_code=0, rows=rows)
                print(f"[validation] PASS {item.name}: {rows:,} rows", flush=True)
            publisher.signal()
    return failures


def inspect_outputs():
    """Read-only validation used by preflight checks."""
    baseline = load_baseline()
    failures = []
    for lane, items in LANES.items():
        for item in items:
            try:
                rows, headers = workbook_metadata(item.output)
                missing = sorted(REQUIRED_HEADERS - headers)
                prior = baseline.get(item.name)
                if rows <= 0:
                    raise ValueError("workbook contains no data rows")
                if missing:
                    raise ValueError("missing required columns: " + ", ".join(missing))
                if prior and rows < int(prior * (1 - DROP_TOLERANCE)):
                    raise ValueError(f"count drop exceeds 5%: {prior:,} -> {rows:,}")
                print(f"PREFLIGHT PASS {item.name}: {rows:,} rows")
            except Exception as exc:
                failures.append((item.name, str(exc)))
                print(f"PREFLIGHT FAIL {item.name}: {exc}", file=sys.stderr)
    return failures


def check_config():
    errors = []
    seen = set()
    total = 0
    for lane, items in LANES.items():
        print(f"{lane}: " + " -> ".join(item.name for item in items))
        for item in items:
            total += 1
            if item.name in seen:
                errors.append(f"Duplicate account: {item.name}")
            seen.add(item.name)
            if not (item.directory / item.script).exists():
                errors.append(f"Missing script: {item.directory / item.script}")
    if total != 22:
        errors.append(f"Expected 22 QA accounts, configured {total}.")
    if errors:
        for error in errors:
            print(f"CONFIG ERROR: {error}", file=sys.stderr)
        return 1
    print(f"Configuration valid: {total} accounts across {len(LANES)} lanes.")
    return 0


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--check-config", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--no-publish", action="store_true")
    args = parser.parse_args()

    if args.check_config:
        return check_config()

    store = StatusStore()
    publisher = LivePublisher(enabled=not args.no_publish)
    results = {item.name: 0 for items in LANES.values() for item in items}

    if args.validate_only:
        failures = inspect_outputs()
        return 1 if failures else 0

    run_id = datetime.now().strftime("%Y%m%d-%H%M%S")
    log_dir = LOG_ROOT / run_id
    log_dir.mkdir(parents=True, exist_ok=True)

    store.queue_all()
    publisher.start()
    publisher.publish("Parallel QA queued")

    threads = []
    for index, (lane, items) in enumerate(LANES.items()):
        thread = threading.Thread(
            target=run_lane,
            args=(lane, items, index * LANE_START_STAGGER_SECONDS, log_dir, store, publisher, results),
            name=slug(lane),
        )
        thread.start()
        threads.append(thread)

    for thread in threads:
        thread.join()

    print("[parallel] All lanes finished. Running validation barrier...", flush=True)
    failures = validate_outputs(results, store, publisher)
    publisher.close("Parallel QA validation complete")

    if failures:
        print(f"[parallel] Validation failed for {len(failures)} account(s). Dashboard build blocked.", flush=True)
        return 1
    print("[parallel] Validation barrier passed for all 22 QA accounts.", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
